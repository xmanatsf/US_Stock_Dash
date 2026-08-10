"""Workbook loading: xlsx -> typed price/volume/mapping frames.

Ported from build_renmac_dashboard.py:60 `load()` (the superset of build_dashboard.py:43)
and build_sp100_dashboard.py:63 `load_mapping()`.

HARD RULE, do not regress: there are NO module-level path globals here. renmac_score.py:125
monkeypatched `build_dashboard.XLSX` at import time to redirect input, which makes the module
unusable for more than one universe in a single process -- and build_all.py runs four universes
in one process. Every path arrives as a function argument sourced from config.
"""

from __future__ import annotations

import datetime as _dt
import os
import re
from dataclasses import dataclass, field

import openpyxl

# ---------------------------------------------------------------- input resolution

_DATE_STAMPED = re.compile(r"^(?P<stem>.+?)[ _](?P<date>\d{8})\.xlsx$", re.IGNORECASE)
_LOCK = re.compile(r"^~\$")


class InputNotFound(Exception):
    pass


class AmbiguousInput(Exception):
    pass


def resolve_input(raw_dir: str, stem: str, date_stamped: bool = True) -> str:
    """Return the newest workbook matching an ANCHORED stem.

    The match is `<exact stem> <YYYYMMDD>.xlsx`, never `*stem*`. This matters: data/raw holds
    both `US software stocks 5y price and volume 20260809.xlsx` and the older, differently
    shaped `US software stocks 20260804.xlsx`. A loose `*software*` glob matches both and picks
    by mtime, which silently selects the wrong file.

    Excel lock files (`~$...`) are excluded -- they are present whenever the user has a workbook
    open, and openpyxl cannot read them.
    """
    if not os.path.isdir(raw_dir):
        raise InputNotFound(f"raw dir does not exist: {raw_dir}")

    stem_l = stem.lower()
    hits: list[tuple[str, str]] = []  # (datestamp, path)

    for name in os.listdir(raw_dir):
        if _LOCK.match(name) or not name.lower().endswith(".xlsx"):
            continue
        if date_stamped:
            m = _DATE_STAMPED.match(name)
            if not m or m.group("stem").lower() != stem_l:
                continue
            hits.append((m.group("date"), os.path.join(raw_dir, name)))
        else:
            if name.lower() != f"{stem_l}.xlsx":
                continue
            hits.append(("", os.path.join(raw_dir, name)))

    if not hits:
        raise InputNotFound(
            f"no workbook matching stem {stem!r} (date_stamped={date_stamped}) in {raw_dir}"
        )
    hits.sort(key=lambda h: h[0], reverse=True)
    if len(hits) > 1 and hits[0][0] == hits[1][0]:
        raise AmbiguousInput(f"two workbooks share the newest datestamp: {hits[0][1]}, {hits[1][1]}")
    return hits[0][1]


# ---------------------------------------------------------------- cell coercion

_PLACEHOLDERS = {"NA", "NM", "#PEND", "SPGRANGEV", "#N/A", "#VALUE!", "#REF!", ""}

# Tokens that mean "CapIQ never resolved this formula". A blank or whitespace-only header label
# is NOT one of these -- it is a cosmetic vendor artifact (CTVA ships a ' ' price label in the
# 2026-08-09 SP500 export while carrying 1,255 clean prints). Treating blank as fatal would
# block a perfectly good universe.
_UNRESOLVED_TOKENS = {"#PEND", "SPGRANGEV", "#N/A", "#VALUE!", "#REF!", "NA", "NM"}


def num(v):
    """CapIQ writes 'NA'/'NM'/'#PEND' as literal strings into otherwise-numeric columns.

    Anything that is not a real int/float becomes None. `bool` is explicitly rejected because
    `isinstance(True, int)` is True in Python, and NaN is rejected because it poisons every
    comparison downstream. The original (build_dashboard.py:38) did a bare isinstance check;
    this is the tightened version.
    """
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if v != v else v  # NaN check
    return None


def _is_placeholder(v) -> bool:
    return isinstance(v, str) and v.strip().upper() in _PLACEHOLDERS


# ---------------------------------------------------------------- date coercion


def to_date(v):
    """Column A normally arrives as a real datetime. Guard the two documented exceptions."""
    if isinstance(v, _dt.datetime):
        return v.date()
    if isinstance(v, _dt.date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        # Excel serial date; 1899-12-30 origin absorbs the Lotus leap-year bug.
        return (_dt.datetime(1899, 12, 30) + _dt.timedelta(days=float(v))).date()
    return None


# ---------------------------------------------------------------- price/volume book


@dataclass
class PriceBook:
    path: str
    dates: list[str] = field(default_factory=list)          # 'YYYY-MM-DD', trading days only
    tickers: list[str] = field(default_factory=list)
    px: dict[str, list] = field(default_factory=dict)
    vol: dict[str, list] = field(default_factory=dict)
    labels: dict[str, str] = field(default_factory=dict)     # row-2 header text per sheet
    calendar: dict = field(default_factory=dict)             # drop report
    coercions: dict[str, int] = field(default_factory=dict)  # ticker -> placeholder count
    raw_dates: list = field(default_factory=list)            # pre-drop, for the QC report

    @property
    def n(self) -> int:
        return len(self.dates)


def _grab(ws, spacer_drop: bool = True):
    """Split a sheet into (ticker header, row-2 label header, data rows).

    THE two-header-row gotcha, and the single most damaging thing to regress:
    row 0 is the ticker header UNCONDITIONALLY, and the `r[0] is not None` filter is applied
    ONLY to rows[1:]. Both header rows carry None in column A, so filtering before splitting
    eats the header and every downstream "ticker" becomes a float.
    """
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], [], []

    header = rows[0]
    labels = rows[1] if len(rows) > 1 else ()
    data = [r for r in rows[1:] if r and r[0] is not None]

    if spacer_drop:
        keep = [j for j in range(1, len(header)) if header[j] is not None]
    else:
        keep = list(range(1, len(header)))

    hdr = [str(header[j]) for j in keep]
    lbl = [str(labels[j]) if j < len(labels) and labels[j] is not None else "" for j in keep]
    trimmed = [(r[0],) + tuple(r[j] if j < len(r) else None for j in keep) for r in data]
    return hdr, lbl, trimmed


def load_price_volume(
    path: str,
    *,
    price_sheet: str = "Price",
    volume_sheet: str = "Volume",
    price_label: str = "Day Close Price",
    volume_label: str = "Volume",
    strict_labels: bool = True,
) -> tuple[PriceBook, list]:
    """Load a two-sheet price/volume workbook into a PriceBook.

    Returns (book, findings). Findings are plain dicts here; validators.py wraps them into
    Finding records. This function does NOT raise on bad data -- a corrupt workbook must still
    load far enough for the validator to describe what is wrong with it, otherwise a tab can
    never render its own validation banner.

    Sheets are read BY NAME, never by index: an add-in cache sheet (`__snloffice`) sits ahead
    of the data sheets in every one of these workbooks.
    """
    findings: list[dict] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sn in (price_sheet, volume_sheet):
            if sn not in wb.sheetnames:
                findings.append(
                    {"code": "SHEET_MISSING", "severity": "fatal",
                     "message": f"sheet {sn!r} not in {os.path.basename(path)}; found {wb.sheetnames}"}
                )
                return PriceBook(path=path), findings

        phdr, plbl, prows = _grab(wb[price_sheet])
        vhdr, vlbl, vrows = _grab(wb[volume_sheet])
    finally:
        wb.close()

    if phdr != vhdr:
        only_p = [t for t in phdr if t not in set(vhdr)]
        only_v = [t for t in vhdr if t not in set(phdr)]
        findings.append(
            {"code": "HEADER_MISMATCH", "severity": "fatal",
             "message": "Price and Volume ticker columns differ",
             "detail": {"priceOnly": only_p[:20], "volumeOnly": only_v[:20]}}
        )

    # Row-2 label check. This is what caught the corrupt HW export ('#PEND'/'SPGRANGEV').
    # Severity splits three ways so a cosmetic blank cannot impersonate an unresolved formula.
    for sheet, hdr, lbl, expect in (
        (price_sheet, phdr, plbl, price_label),
        (volume_sheet, vhdr, vlbl, volume_label),
    ):
        unresolved, blank, other = [], [], []
        for t, x in zip(hdr, lbl):
            if x == expect:
                continue
            if x.strip().upper() in _UNRESOLVED_TOKENS:
                unresolved.append(t)
            elif not x.strip():
                blank.append(t)
            else:
                other.append((t, x))
        if unresolved:
            findings.append(
                {"code": "HEADER_PLACEHOLDER", "severity": "fatal" if strict_labels else "warn",
                 "message": f"{sheet} row-2 label is an unresolved CapIQ formula token "
                            f"('#PEND'/'SPGRANGEV') on {len(unresolved)} column(s); the export "
                            f"almost certainly saved a previous query's cached values",
                 "tickers": sorted(unresolved)[:40], "count": len(unresolved)}
            )
        if blank:
            findings.append(
                {"code": "HEADER_LABEL_BLANK", "severity": "info",
                 "message": f"{sheet} row-2 label is blank on {len(blank)} column(s) "
                            f"(cosmetic vendor artifact; data itself is checked separately)",
                 "tickers": sorted(blank)[:40], "count": len(blank)}
            )
        if other:
            findings.append(
                {"code": "HEADER_LABEL_UNEXPECTED", "severity": "warn",
                 "message": f"{sheet} row-2 label is unexpected on {len(other)} column(s), "
                            f"expected {expect!r}",
                 "tickers": sorted(t for t, _ in other)[:40], "count": len(other),
                 "detail": {"samples": other[:5]}}
            )

    # Join Volume to Price BY DATE, never by row position -- a vendor-side range mismatch
    # would otherwise shift every volume reading against the wrong close.
    vby = {}
    for r in vrows:
        d = to_date(r[0])
        if d is not None:
            vby[d] = r

    dates, coerc = [], {}
    px_rows, vol_rows = [], []
    missing_vol = []
    for r in prows:
        d = to_date(r[0])
        if d is None:
            continue
        dates.append(d)
        px_rows.append(r)
        vr = vby.get(d)
        if vr is None:
            missing_vol.append(d)
            vol_rows.append((r[0],) + (None,) * (len(phdr)))
        else:
            vol_rows.append(vr)

    if missing_vol:
        findings.append(
            {"code": "PXVOL_DATE_MISMATCH", "severity": "warn",
             "message": f"{len(missing_vol)} Price dates have no matching Volume row",
             "count": len(missing_vol),
             "detail": {"first": str(missing_vol[0]), "last": str(missing_vol[-1])}}
        )

    # dupes / monotonicity are reported by validators.check_date_grid; build the book here.
    book = PriceBook(path=path, tickers=list(phdr))
    book.labels = {price_sheet: plbl[0] if plbl else "", volume_sheet: vlbl[0] if vlbl else ""}
    book.raw_dates = list(dates)
    book.dates = [d.isoformat() for d in dates]

    for j, t in enumerate(phdr, start=1):
        pc, vc, bad = [], [], 0
        for i in range(len(px_rows)):
            rawp = px_rows[i][j] if j < len(px_rows[i]) else None
            rawv = vol_rows[i][j] if j < len(vol_rows[i]) else None
            if _is_placeholder(rawp):
                bad += 1
            pc.append(num(rawp))
            vc.append(num(rawv))
        book.px[t] = pc
        book.vol[t] = vc
        if bad:
            coerc[t] = bad
    book.coercions = coerc

    dup = [t for t in phdr if phdr.count(t) > 1]
    if dup:
        findings.append(
            {"code": "DUPLICATE_TICKER", "severity": "warn",
             "message": "duplicate ticker columns", "tickers": sorted(set(dup))}
        )

    return book, findings


# ---------------------------------------------------------------- calendar-day removal


def drop_calendar_days(book: PriceBook, threshold: float = 0.95) -> tuple[PriceBook, dict]:
    """Reduce a calendar-day grid to real trading days. ORDER IS LOAD-BEARING.

    Stage 1 -- weekends: drop Sat/Sun rows. These carry the prior close AND the prior volume
    forward-filled, so a '20-dma' left uncorrected spans ~14 real sessions and every weekend
    counts Friday's volume three times. (1,827 -> 1,305 on all four workbooks.)

    Stage 2 -- holidays: drop any weekday row that repeats the PRIOR SURVIVING row across
    >threshold of the universe on price AND volume. (1,305 -> 1,255.)

    The price-AND-volume test is not optional. Exact whole-row equality under-detects badly --
    it finds 11 holidays in the software workbook instead of 50 -- because a handful of names
    carry a differing placeholder. Comparing only tickers where both rows have real numbers,
    and requiring >95% agreement, recovers exactly the 50 US market holidays.
    """
    n_raw = len(book.dates)
    tickers = book.tickers
    keep_idx: list[int] = []
    weekend, holidays = [], []

    for i, d in enumerate(book.raw_dates):
        if d.weekday() >= 5:
            weekend.append(d)
            continue
        if keep_idx:
            p = keep_idx[-1]
            same = tot = 0
            for t in tickers:
                a, b = book.px[t][i], book.px[t][p]
                c, e = book.vol[t][i], book.vol[t][p]
                if a is None or b is None:
                    continue
                tot += 1
                if a == b and c == e:
                    same += 1
            if tot and same / tot > threshold:
                holidays.append(d)
                continue
        keep_idx.append(i)

    out = PriceBook(path=book.path, tickers=tickers, labels=book.labels, coercions=book.coercions)
    out.raw_dates = [book.raw_dates[i] for i in keep_idx]
    out.dates = [d.isoformat() for d in out.raw_dates]
    out.px = {t: [book.px[t][i] for i in keep_idx] for t in tickers}
    out.vol = {t: [book.vol[t][i] for i in keep_idx] for t in tickers}

    years = ((out.raw_dates[-1] - out.raw_dates[0]).days / 365.25) if len(out.raw_dates) > 1 else 0
    report = {
        "rawRows": n_raw,
        "weekendRows": len(weekend),
        "afterWeekendDrop": n_raw - len(weekend),
        "holidayRows": len(holidays),
        "tradingDays": len(out.dates),
        "sessionsPerYear": round(len(out.dates) / years, 1) if years else None,
        "firstDate": out.dates[0] if out.dates else None,
        "lastDate": out.dates[-1] if out.dates else None,
        "holidayDates": [d.isoformat() for d in holidays],
    }
    out.calendar = report
    return out, report


def first_real_index(col) -> int | None:
    """Index of the first non-None value, i.e. the ticker's first real print.

    Every 'NA' run in these workbooks is leading-edge only (verified: zero interior, zero
    trailing across all four price files), so this is the listing date, not a gap.
    """
    for i, v in enumerate(col):
        if v is not None:
            return i
    return None


def ffill_small_gaps(col, max_gap: int = 3):
    """Bounded forward-fill of PRICE gaps only (build_dashboard.py:100).

    Volume is never forward-filled -- a filled volume print fabricates activity that did not
    happen, and zero is real data (a halt, or a thin name that did not trade).

    On the current vintages this is a no-op, since all NA runs are leading. It is retained as a
    defensive validator: if it ever reports a non-zero fill count, an interior gap has appeared
    and the 'leading-only' assumption needs re-checking.
    """
    out = list(col)
    filled = 0
    i = 0
    n = len(out)
    while i < n:
        if out[i] is not None:
            i += 1
            continue
        j = i
        while j < n and out[j] is None:
            j += 1
        prev = out[i - 1] if i > 0 else None
        if prev is not None and j < n and (j - i) <= max_gap:
            for k in range(i, j):
                out[k] = prev
            filled += j - i
        i = j
    return out, filled


# ---------------------------------------------------------------- mapping sheets


def load_mapping(
    path: str,
    *,
    sheet: str,
    data_start_row: int,
    columns: dict,
) -> tuple[dict, list]:
    """Load a CapIQ mapping export into {ticker: {field: value}}.

    `data_start_row` and `columns` are 1-based and come from sectors.json, never hardcoded.
    The two mapping files in this project have DIFFERENT preambles and different column orders
    (SP500 industry groups: ticker at col 3, data from row 6; US network and hw stocks: ticker
    at col 4, market cap at col 3), and build_sp100_dashboard.py:68 uses yet another offset for
    SP100.xlsx. An off-by-two here silently files every ticker under the wrong sector, and the
    resulting breadth rollups look plausible while being nonsense.
    """
    findings: list[dict] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            findings.append(
                {"code": "SHEET_MISSING", "severity": "fatal",
                 "message": f"sheet {sheet!r} not in {os.path.basename(path)}; found {wb.sheetnames}"}
            )
            return {}, findings
        rows = list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()

    tcol = columns["ticker"]
    out: dict = {}
    for r in rows[data_start_row - 1:]:
        if not r or len(r) < tcol or r[tcol - 1] is None:
            continue
        rec = {}
        for fname, cidx in columns.items():
            v = r[cidx - 1] if len(r) >= cidx else None
            rec[fname] = str(v).strip() if isinstance(v, str) else v
        out[str(r[tcol - 1]).strip()] = rec

    if not out:
        findings.append(
            {"code": "MAPPING_EMPTY", "severity": "fatal",
             "message": f"no rows parsed from {os.path.basename(path)} at dataStartRow={data_start_row}; "
                        f"print the first 8 rows before assuming the header offset"}
        )
    return out, findings
